"""Step 7 (sheet tabs + History) and Step 11 (monthly history archive). Reads the validated build; never recalculates."""
import csv, json, shutil, sys

import config

sys.path.insert(0, str(config.SCRIPTS))
import build_report as br  # noqa: E402  (reused for the approved cleaning rules only: DUP_KEY, line_id)

HIST_COLS = ["month", "report", "sku", "sku_type", "qty_sold", "sales_amount", "amazon_units", "ebay_units", "bq_units",
             "total_units", "avg_daily_units", "run_id"]


def tabs(month, data):
    """The 7 tabs for the reporting month (History is added by the caller)."""
    cat = data.get("sku_cat", {})
    t = {}
    for mp in config.MARKETPLACES:
        t[mp] = [["Month", "SKU (as sold)", "SKU Type", "Quantity Sold", "Sales Amount (GBP)", "Avg Price", "Qty MoM %", "Sales MoM %"]] + \
                [[r[0], r[2], r[3], r[5], r[6], r[7], r[10], r[11]] for r in data["r13"] if r[0] == month and r[1] == mp]
    t["Single SKU Units"] = [["Month", "Single SKU", "Product Category", "Amazon Units", "eBay Units", "B&Q Units", "Total Units", "Avg Daily Units"]] + \
        [[r[0], r[1], cat.get(r[1], ["Category Not Available"])[0], r[2], r[3], r[4], r[5], r[6]] for r in data["r4"] if r[0] == month]
    t["Combo Components"] = [["Month", "Marketplace", "Type", "Combo / Pack SKU", "Component SKU", "Qty Sold", "Qty per unit", "Units Added"]] + \
        [[r[0], r[1], "Combo", r[2], r[3], r[4], (r[5] // r[4]) if r[4] else None, r[5]] for r in data["combo"] if r[0] == month] + \
        [[r[0], r[1], "Pack", r[2], r[3], r[4], (r[5] // r[4]) if r[4] else None, r[5]] for r in data["pack"] if r[0] == month]
    t["Exceptions"] = [["Marketplace", "ID", "Qty", "Sales (GBP)", "Month", "Reason", "Details", "Impact", "Status"]] + \
        [[e["mp"], e["id"], e["qty"], e["sales"], e["month"], e["type"], e["details"], e["impact"], e["status"]] for e in data["exc"] if e["month"] == month]
    return t


def history_rows(month, data, run_id):
    rows = [dict(month=r[0], report=r[1], sku=r[2], sku_type=r[3], qty_sold=r[5], sales_amount=r[6], amazon_units="", ebay_units="",
                 bq_units="", total_units="", avg_daily_units="", run_id=run_id) for r in data["r13"] if r[0] == month]
    rows += [dict(month=r[0], report="Single SKU Units", sku=r[1], sku_type="", qty_sold="", sales_amount="", amazon_units=r[2],
                  ebay_units=r[3], bq_units=r[4], total_units=r[5], avg_daily_units=r[6], run_id=run_id) for r in data["r4"] if r[0] == month]
    return [{k: str(v) if v is not None else "" for k, v in r.items()} for r in rows]


def append_history(month, data, run_id):
    """Append-only. A month already present is never rewritten; a re-run that differs is reported, not applied."""
    new = history_rows(month, data, run_id)
    existing = []
    if config.HISTORY_MASTER.exists():
        with config.HISTORY_MASTER.open(newline="", encoding="utf-8") as fh:
            existing = list(csv.DictReader(fh))
    have = [r for r in existing if r["month"] == month]
    if have:
        strip = lambda rows: sorted(json.dumps({k: v for k, v in r.items() if k != "run_id"}, sort_keys=True) for r in rows)
        same = strip(have) == strip(new)
        return ("ALREADY PRESENT - identical" if same else "ALREADY PRESENT - DIFFERENT (not overwritten)"), same
    config.HISTORY_MASTER.parent.mkdir(parents=True, exist_ok=True)
    first = not config.HISTORY_MASTER.exists()
    with config.HISTORY_MASTER.open("a", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=HIST_COLS)
        if first:
            w.writeheader()
        w.writerows(new)
    return f"APPENDED {len(new)} rows", True


def write_workbook(path, month, data):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); wb.remove(wb.active)
    t = tabs(month, data)
    hist = []
    if config.HISTORY_MASTER.exists():
        with config.HISTORY_MASTER.open(newline="", encoding="utf-8") as fh:
            hist = list(csv.reader(fh))
    t["History"] = hist or [HIST_COLS]
    for name in config.SHEET_TABS:
        ws = wb.create_sheet(name[:31])
        for row in t[name]:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions  # Month / Marketplace / SKU filters
    wb.save(path)
    return {k: len(v) - 1 for k, v in t.items()}


def push_google_sheet(month, data):
    """Appends the month to each tab of the configured spreadsheet (months already present are skipped)."""
    sid = config.GOOGLE_SHEET["spreadsheet_id"]
    if not sid:
        return "NOT CONFIGURED", "No Product Performance spreadsheet id set (see config.GOOGLE_SHEET)"
    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build
    creds = Credentials.from_service_account_file(config.GOOGLE_SHEET["credentials_file"], scopes=["https://www.googleapis.com/auth/spreadsheets"])
    svc = build("sheets", "v4", credentials=creds, cache_discovery=False).spreadsheets()
    meta = svc.get(spreadsheetId=sid).execute()
    have = {s["properties"]["title"] for s in meta["sheets"]}
    add = [{"addSheet": {"properties": {"title": n}}} for n in config.SHEET_TABS if n not in have]
    if add:
        svc.batchUpdate(spreadsheetId=sid, body={"requests": add}).execute()
    t = tabs(month, data)
    t["History"] = [HIST_COLS] + [[r[c] for c in HIST_COLS] for r in history_rows(month, data, "sheet")]
    done = []
    for name in config.SHEET_TABS:
        col = svc.values().get(spreadsheetId=sid, range=f"'{name}'!A:E").execute().get("values", [])
        month_col = 4 if name == "Exceptions" else 0
        if any(len(r) > month_col and r[month_col] == month for r in col[1:]):
            done.append(f"{name}: month already present (skipped)")
            continue
        rows = t[name] if not col else t[name][1:]
        svc.values().append(spreadsheetId=sid, range=f"'{name}'!A1", valueInputOption="RAW", insertDataOption="INSERT_ROWS",
                            body={"values": [[("" if v is None else v) for v in r] for r in rows]}).execute()
        done.append(f"{name}: +{len(rows)}")
    return "UPDATED", "; ".join(done)


def archive(month, run_id, data, html_path, workbook_path, log_record):
    """14_History/<month>/run_<run_id>/ - a new folder every run, nothing overwritten."""
    d = config.HISTORY_DIR / month / f"run_{run_id}"
    (d / "raw").mkdir(parents=True, exist_ok=False)
    for mp, f in config.RAW_FILES.items():
        src = config.RAW_DIR / month / f
        if src.exists():
            shutil.copy2(src, d / "raw" / f)
    # cleaned lines: approved cleaning only (text IDs, trim, ASIN uppercase, exact-duplicate flag) - mapping is in the reports
    (d / "cleaned").mkdir()
    for mp, f in config.RAW_FILES.items():
        src = config.RAW_DIR / month / f
        if not src.exists():
            continue
        with src.open(newline="", encoding="utf-8") as fh:
            rows = list(csv.DictReader(fh))
        seen = set()
        with (d / "cleaned" / f"Clean_{f.split('_', 1)[1]}").open("w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["order_item_info", "marketplace", "id_clean", "sku", "quantity", "sales_amount", "order_date", "exact_duplicate"])
            for r in rows:
                key = tuple((r[k] or "").strip() for k in br.DUP_KEY)
                w.writerow([r["order_item_info"], mp, br.line_id(mp, r), (r["sku"] or "").strip(), r["quantity"], r["order_total"],
                            r["order_date"], "yes" if key in seen else "no"])
                seen.add(key)
    shutil.copy2(html_path, d / html_path.name)
    shutil.copy2(workbook_path, d / workbook_path.name)
    (d / f"data_{month}.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    t = tabs(month, data)
    for name in ("Amazon", "eBay", "B&Q", "Single SKU Units", "Exceptions"):
        with (d / f"{name.replace('&', 'and').replace(' ', '_')}.csv").open("w", newline="", encoding="utf-8") as fh:
            csv.writer(fh).writerows(t[name])
    with (d / "reconciliation.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh); w.writerow(["check", "month", "marketplace", "status", "detail"])
        for k in ("recon13", "recon4", "recon_ph", "recon_cat"):
            for r in data.get(k, []):
                if r["month"] == month:
                    w.writerow([k, r["month"], r.get("mp", "Amazon" if k == "recon_ph" else "all"), r["status"], json.dumps(r)])
    (d / "run_log.json").write_text(json.dumps(log_record, indent=1, default=str), encoding="utf-8")
    return d
